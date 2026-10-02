import openpyxl, random, os, sys
from collections import Counter

inp = sys.argv[1] if len(sys.argv) > 1 else 'DSB_Input_Template.xlsx'
out = inp

wb = openpyxl.load_workbook(inp)
ws = wb['Players']
bws = wb['Blacklist']

# Guarantee Players.Status exists in the workbook before any processing.
# If it already exists, reuse it; otherwise append it as the last column.
STATUS_HEADER = 'Status'
status_col = None
for _c in range(1, ws.max_column + 1):
    if str(ws.cell(1, _c).value or '').strip().casefold() == STATUS_HEADER.casefold():
        status_col = _c
        break
if status_col is None:
    status_col = ws.max_column + 1
    ws.cell(1, status_col).value = STATUS_HEADER

cws = wb['Config'] if 'Config' in wb.sheetnames else None

def normalize_name(value):
    return ' '.join(str(value or '').strip().split()).casefold()

# -----------------------------
# Config
# -----------------------------
cfg = {}
if cws is not None:
    for r in cws.iter_rows(min_row=2, values_only=True):
        if r[0] is not None:
            cfg[str(r[0]).strip()] = r[1]

mode = str(cfg.get('Mode', 'SAVING')).upper()
seed = int(cfg.get('Random Seed', 42))

MAX_TEAM = int(cfg.get('Max Team Size', 30))
MAX_MAIN = int(cfg.get('Main Players per Team', 20))
MAX_RESERVE = int(cfg.get('Reserve Players per Team', 10))

# Number of players taken from the persistent DSB Priority queue on each run.
DSB_PRIORITY_COUNT = int(cfg.get('DSB Priority Count', 0))

rng = random.Random(seed)

# -----------------------------
# Blacklist
# -----------------------------
black = {
    normalize_name(r[0])
    for r in bws.iter_rows(min_row=2, values_only=True)
    if r[0] and str(r[1] or '').strip().casefold() == 'blacklist dsb next week'
}

# -----------------------------
# Players / eligible pool
# -----------------------------
headers = [c.value for c in ws[1]]
idx = {h: i for i, h in enumerate(headers)}

players = []
for r in ws.iter_rows(min_row=2, values_only=True):
    name = r[idx['Name']]
    if not name:
        continue

    # Blacklist is always applied before every selection stage.
    if normalize_name(name) in black:
        continue

    av = str(r[idx['Team Availability']] or '')
    thp = r[idx['THP']]
    dav = r[idx['DA/VS']]

    players.append({
        'Name': name,
        'Availability': av,
        'THP': float(thp) if isinstance(thp, (int, float)) else 0.0,
        'DA/VS': float(dav) if isinstance(dav, (int, float)) else 0.0,
    })

eligible_by_key = {normalize_name(p['Name']): p for p in players}
eligible = list(players)

# -----------------------------
# DSB Priority persistent queue
# -----------------------------
if 'DSB Priority' not in wb.sheetnames:
    pws = wb.create_sheet('DSB Priority')
    pws.append(['Name'])
else:
    pws = wb['DSB Priority']

# Read existing queue in sheet order.
priority_names = []
seen_priority = set()
for r in pws.iter_rows(min_row=2, values_only=True):
    raw = r[0] if r else None
    key = normalize_name(raw)
    if not key or key in seen_priority:
        continue
    # Never retain blacklisted or no-longer-eligible players in the queue.
    if key in black or key not in eligible_by_key:
        continue
    seen_priority.add(key)
    priority_names.append(key)

# -----------------------------
# Selection priority:
# 1. Top 10 THP
# 2. Top 40 DA/VS (ALL-IN only)
# 3. DSB Priority queue, X players
# 4. Spin
# -----------------------------
eligible_sorted = sorted(eligible, key=lambda p: (p['THP'], p['Name']), reverse=True)
top10 = eligible_sorted[:min(10, len(eligible_sorted))]

selected = {}
reason = {}

def add_selection(p, why):
    key = normalize_name(p['Name'])
    if key not in selected:
        selected[key] = p
        reason[key] = why
        return True
    return False

for p in top10:
    add_selection(p, 'Top 10 THP')

# Top 40 DA/VS applies only in ALL-IN.
da40 = []
if mode == 'ALL-IN':
    remaining_for_da = [p for p in eligible if normalize_name(p['Name']) not in selected]
    da40 = sorted(
        remaining_for_da,
        key=lambda p: (p['DA/VS'], p['THP'], p['Name']),
        reverse=True
    )[:min(40, len(remaining_for_da))]
    for p in da40:
        add_selection(p, 'Top 40 DA/VS')

# DSB Priority: take X AVAILABLE players from the existing queue.
# Availability is always taken from the current Players sheet. A priority
# member with blank/invalid availability is skipped and the next available
# priority member is considered.
def is_currently_available(p):
    av = str(p.get('Availability', '') or '').strip().lower()
    return av in ('a', 'b', 'a+b', 'a only', 'b only')

# Select up to X *additional* valid priority members. Walk the queue rather
# than slicing it so every valid candidate is considered until the requested
# count is reached. A member already selected by a higher-priority stage is
# skipped but does not consume the DSB Priority Count.
priority_take = []
priority_limit = max(0, DSB_PRIORITY_COUNT)
for k in priority_names:
    if len(priority_take) >= priority_limit:
        break
    p = eligible_by_key.get(k)
    if p is None:
        continue
    if k in selected:
        continue
    if not is_currently_available(p):
        continue
    priority_take.append(p)
    add_selection(p, 'DSB Priority')

# -----------------------------
# Availability helpers
# -----------------------------
def av_lower(p):
    return p['Availability'].strip().lower()

def can_a(p):
    a = av_lower(p)
    return a in ('a', 'a only') or 'a+b' in a

def can_b(p):
    a = av_lower(p)
    return a in ('b', 'b only') or 'a+b' in a

def is_a_only(p):
    return av_lower(p) in ('a', 'a only')

def is_b_only(p):
    return av_lower(p) in ('b', 'b only')

def is_flex(p):
    a = av_lower(p)
    return 'a+b' in a

A_pool = [p for p in eligible if can_a(p)]
B_pool = [p for p in eligible if can_b(p)]

# Maximum feasible total under hard team limits and availability.
target_total = min(
    len(eligible),
    MAX_TEAM * 2,
    MAX_TEAM + len(B_pool),
    MAX_TEAM + len(A_pool),
)

# The selected priority stages may include an availability combination that
# cannot fit in a valid A/B split. Never exceed hard team limits.
# We first preserve all high-priority selections, then use Spin to fill.
def feasible_split_for(chosen):
    n = len(chosen)
    a_only = sum(is_a_only(p) for p in chosen)
    b_only = sum(is_b_only(p) for p in chosen)
    flex = sum(is_flex(p) for p in chosen)
    lo = a_only
    hi = a_only + flex
    return [
        a for a in range(lo, hi + 1)
        if a <= MAX_TEAM and n - a <= MAX_TEAM
    ]

# Enforce the hard roster capacity while preserving selection priority:
# Top 10 THP > Top 40 DA/VS > DSB Priority > Spin.
# If a constrained availability mix makes the current set impossible, remove
# lower-priority players first, then refill with Spin.
priority_rank = {
    'Top 10 THP': 1,
    'Top 40 DA/VS': 2,
    'DSB Priority': 3,
    'Spin': 4,
}

while len(selected) > target_total or not feasible_split_for(list(selected.values())):
    removable = [
        (k, reason[k]) for k in selected
        if reason[k] != 'Top 10 THP'
    ]
    if not removable:
        break
    # Remove the lowest selection priority first; within the same stage,
    # remove the last-added member so DSB Priority preserves queue order.
    max_rank = max(priority_rank[r] for _, r in removable)
    candidates = [k for k, r in removable if priority_rank[r] == max_rank]
    k = candidates[-1]
    del selected[k]
    del reason[k]

# Spin fills remaining capacity while reserving the required A/B capacity.
# This prevents availability-heavy random selection from creating an invalid
# roster (for example, more than 30 A players).
n_target = target_total

# Choose the most balanced feasible final split for the full eligible pool.
possible_splits = []
for a_size in range(MAX_TEAM + 1):
    b_size = n_target - a_size
    if b_size < 0 or b_size > MAX_TEAM:
        continue
    if a_size > len(A_pool) or b_size > len(B_pool):
        continue
    possible_splits.append((a_size, b_size))

if not possible_splits:
    raise RuntimeError('No feasible target team split.')

targetA, targetB = min(
    possible_splits,
    key=lambda x: (abs(x[0] - x[1]), -x[0])
)

selected_values = list(selected.values())
sel_a_only = sum(is_a_only(p) for p in selected_values)
sel_b_only = sum(is_b_only(p) for p in selected_values)
sel_flex = sum(is_flex(p) for p in selected_values)

# If the higher-priority selections cannot fit the target split, remove the
# lowest-priority stage until they do.
def selected_split_possible():
    for f_a in range(sel_flex + 1):
        f_b = sel_flex - f_a
        if sel_a_only + f_a <= targetA and sel_b_only + f_b <= targetB:
            return True
    return False

while (
    len(selected) > target_total
    or sel_a_only > targetA
    or sel_b_only > targetB
    or not selected_split_possible()
):
    removable = [
        k for k in selected
        if reason[k] != 'Top 10 THP'
    ]
    if not removable:
        break

    max_rank = max(priority_rank[reason[k]] for k in removable)
    candidates = [k for k in removable if priority_rank[reason[k]] == max_rank]
    k = candidates[-1]

    del selected[k]
    del reason[k]

    selected_values = list(selected.values())
    sel_a_only = sum(is_a_only(p) for p in selected_values)
    sel_b_only = sum(is_b_only(p) for p in selected_values)
    sel_flex = sum(is_flex(p) for p in selected_values)

# Recalculate the feasible assignment of already-selected flex players.
flex_assignment_options = []
for f_a in range(sel_flex + 1):
    f_b = sel_flex - f_a
    need_a = targetA - sel_a_only - f_a
    need_b = targetB - sel_b_only - f_b
    if need_a >= 0 and need_b >= 0:
        flex_assignment_options.append((f_a, f_b, need_a, need_b))

if not flex_assignment_options:
    raise RuntimeError('Higher-priority selections cannot fit the target split.')

# Prefer the option requiring the least constrained spin, then balance.
f_a, f_b, need_a, need_b = min(
    flex_assignment_options,
    key=lambda x: (max(x[2], x[3]), abs(x[2] - x[3]))
)

remaining = [
    p for p in eligible
    if normalize_name(p['Name']) not in selected
]

b_only_remaining = [p for p in remaining if is_b_only(p)]
a_only_remaining = [p for p in remaining if is_a_only(p)]
flex_remaining = [p for p in remaining if is_flex(p)]

rng.shuffle(b_only_remaining)
rng.shuffle(a_only_remaining)
rng.shuffle(flex_remaining)

# Fill exact B-only and A-only requirements first.
for p in b_only_remaining:
    if need_b <= 0:
        break
    add_selection(p, 'Spin')
    need_b -= 1

for p in a_only_remaining:
    if need_a <= 0:
        break
    add_selection(p, 'Spin')
    need_a -= 1

# If fixed-availability players were insufficient, use remaining flex players.
for p in flex_remaining:
    if need_a <= 0 and need_b <= 0:
        break
    if need_b > 0:
        add_selection(p, 'Spin')
        need_b -= 1
    elif need_a > 0:
        add_selection(p, 'Spin')
        need_a -= 1

# Final fallback: if the selected set is still short because of an unusual
# availability mix, use any remaining eligible player as Spin.
if len(selected) < target_total:
    fallback = [
        p for p in eligible
        if normalize_name(p['Name']) not in selected
    ]
    rng.shuffle(fallback)
    for p in fallback:
        if len(selected) >= target_total:
            break
        add_selection(p, 'Spin')

chosen = list(selected.values())

# -----------------------------
# Determine valid A/B split and assign teams
# -----------------------------
splits = feasible_split_for(chosen)
if not splits:
    raise RuntimeError('No valid team split under Config capacity limits.')

n = len(chosen)
preferred_a = min(MAX_TEAM, round(n / 2))
target_a = min(splits, key=lambda a: (abs(a - preferred_a), abs((n-a) - preferred_a)))

a_only = [p for p in chosen if is_a_only(p)]
b_only = [p for p in chosen if is_b_only(p)]
flex = [p for p in chosen if is_flex(p)]

need_flex_a = target_a - len(a_only)
if not 0 <= need_flex_a <= len(flex):
    raise RuntimeError('Invalid flexible-player allocation.')

# Fast team balancing: do NOT enumerate all combinations.
# The previous implementation used itertools.combinations(), which is
# exponential (e.g. C(25,12) is already 5.2 million).
# Instead, greedily assign the strongest flexible players to the currently
# weaker side while respecting the exact number required for Team A, then
# perform local pair swaps until no swap improves the THP difference.
flex_sorted = sorted(
    flex, key=lambda p: (p['THP'], p['Name']), reverse=True
)

A_flex_list = []
B_flex_list = []
pa_fixed = sum(p['THP'] for p in a_only)
pb_fixed = sum(p['THP'] for p in b_only)
pa = pa_fixed
pb = pb_fixed

for p in flex_sorted:
    if len(A_flex_list) >= need_flex_a:
        B_flex_list.append(p)
    elif len(B_flex_list) >= len(flex) - need_flex_a:
        A_flex_list.append(p)
        pa += p['THP']
    elif pa <= pb:
        A_flex_list.append(p)
        pa += p['THP']
    else:
        B_flex_list.append(p)
        pb += p['THP']

# Local improvement: swap one A-flex with one B-flex when it reduces
# the absolute THP difference. This is O(F^2) and remains tiny for DSB.
while True:
    current_diff = abs(pa - pb)
    best_swap = None
    best_diff = current_diff

    for ai, ap in enumerate(A_flex_list):
        for bi, bp in enumerate(B_flex_list):
            new_pa = pa - ap['THP'] + bp['THP']
            new_pb = pb - bp['THP'] + ap['THP']
            diff = abs(new_pa - new_pb)
            if diff < best_diff:
                best_diff = diff
                best_swap = (ai, bi, new_pa, new_pb)

    if best_swap is None:
        break

    ai, bi, pa, pb = best_swap
    A_flex_list[ai], B_flex_list[bi] = B_flex_list[bi], A_flex_list[ai]

A_flex = {normalize_name(p['Name']) for p in A_flex_list}

for p in chosen:
    p['Team'] = 'A' if is_a_only(p) or normalize_name(p['Name']) in A_flex else 'B'
    p['Selection Reason'] = reason[normalize_name(p['Name'])]

# -----------------------------
# Main / Reserve / Tier / Slot
# -----------------------------
pattern = [
    'T1','T2','T3','T4','T4','T3','T2','T1',
    'T1','T2','T3','T4','T4','T3','T2','T1',
    'T1','T2','T3','T4'
]

for team in ['A', 'B']:
    arr = sorted(
        [p for p in chosen if p['Team'] == team],
        key=lambda p: (p['THP'], p['Name']),
        reverse=True
    )

    main_n = min(MAX_MAIN, len(arr))
    for i, p in enumerate(arr):
        p['Status'] = 'Main' if i < main_n else 'Reserve'

    main = [p for p in arr if p['Status'] == 'Main']
    reserves = [p for p in arr if p['Status'] == 'Reserve']

    for p, tier in zip(main, pattern):
        p['Tier'] = tier

    for p in reserves:
        p['Tier'] = 'Reserve'

    for tier in ['T1','T2','T3','T4']:
        tier_players = sorted(
            [p for p in main if p['Tier'] == tier],
            key=lambda p: (p['THP'], p['Name']),
            reverse=True
        )
        for slot, p in enumerate(tier_players, 1):
            p['Slot'] = slot

    for slot, p in enumerate(
        sorted(reserves, key=lambda p: (p['THP'], p['Name']), reverse=True), 1
    ):
        p['Slot'] = slot

# -----------------------------
# Update DSB Priority queue
# -----------------------------
selected_keys = {normalize_name(p['Name']) for p in chosen}

# Remove every selected member from the queue.
remaining_priority = [
    k for k in priority_names
    if k not in selected_keys
]

# Add every eligible member who was NOT selected in this run.
# Preserve existing queue order and never duplicate a member.
queue = list(remaining_priority)
queue_set = set(queue)

for p in eligible:
    k = normalize_name(p['Name'])
    if k not in selected_keys and k not in black and k not in queue_set:
        queue.append(k)
        queue_set.add(k)

# Rewrite DSB Priority as a clean, unique list using canonical current Players names.
canonical_by_key = {normalize_name(p['Name']): p['Name'] for p in eligible}
pws.delete_rows(1, pws.max_row)
pws.append(['Name'])
for k in queue:
    if k in canonical_by_key and k not in black:
        pws.append([canonical_by_key[k]])

# -----------------------------
# Update Config to document new setting
# -----------------------------
if cws is not None:
    config_rows = {str(cws.cell(r, 1).value).strip(): r for r in range(2, cws.max_row + 1)}
    if 'DSB Priority Count' in config_rows:
        cws.cell(config_rows['DSB Priority Count'], 2).value = DSB_PRIORITY_COUNT
    else:
        cws.append(['DSB Priority Count', DSB_PRIORITY_COUNT])

    if 'Selection Reason Values' in config_rows:
        cws.cell(config_rows['Selection Reason Values'], 2).value = (
            'Top 10 THP; Top 40 DA/VS; DSB Priority; Spin'
        )

# -----------------------------
# DSB Assignment output
# -----------------------------
if 'DSB Assignment' in wb.sheetnames:
    del wb['DSB Assignment']

ows = wb.create_sheet('DSB Assignment')
cols = [
    'Name','Team','Status','Tier','Slot',
    'THP','DA/VS','Availability','Selection Reason'
]
ows.append(cols)

def sort_key(p):
    team_order = 0 if p['Team'] == 'A' else 1
    status_order = 0 if p['Status'] == 'Main' else 1
    tier_order = {'T1':1,'T2':2,'T3':3,'T4':4,'Reserve':5}.get(p['Tier'], 9)
    return (team_order, status_order, tier_order, p.get('Slot', 99), -p['THP'])

for p in sorted(chosen, key=sort_key):
    ows.append([
        p['Name'], p['Team'], p['Status'], p['Tier'], p['Slot'],
        p['THP'], p['DA/VS'], p['Availability'], p['Selection Reason']
    ])

# -----------------------------
# Update Players.Status from the FINAL run result
# -----------------------------
selected_reason_by_key = {
    normalize_name(p['Name']): p['Selection Reason']
    for p in chosen
}

for r in range(2, ws.max_row + 1):
    key = normalize_name(ws.cell(r, idx['Name'] + 1).value)

    if not key:
        ws.cell(r, status_col).value = ''
    elif key in black:
        ws.cell(r, status_col).value = 'Blacklisted'
    elif key in selected_reason_by_key:
        ws.cell(r, status_col).value = selected_reason_by_key[key]
    else:
        ws.cell(r, status_col).value = 'Not Selected'

wb.save(out)

print('saved', out)
print('mode', mode, 'seed', seed, 'dsb_priority_count', DSB_PRIORITY_COUNT)
print('eligible', len(eligible), 'selected', len(chosen))
print('team', Counter(p['Team'] for p in chosen))
print('power',
      sum(p['THP'] for p in chosen if p['Team'] == 'A'),
      sum(p['THP'] for p in chosen if p['Team'] == 'B'))
print('priority_selected', [
    p['Name'] for p in chosen if p['Selection Reason'] == 'DSB Priority'
])
print('priority_queue_size', len(queue))
